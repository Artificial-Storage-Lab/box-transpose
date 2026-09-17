#!/usr/bin/env python3
"""Draw each slide of proposal.pptx as a PNG, for layout review.

There is no PowerPoint and no LibreOffice in this container, so the usual
"convert to PDF and look at it" check is not available. This re-reads the
shapes out of the .pptx and draws them with their real positions, sizes and
colours, which is enough to catch the defects that actually happen: text
overflowing its box, boxes overlapping, and anything sitting past the margin.

It is an approximation, not a renderer. Fonts are DejaVu rather than
Calibri/Cambria, and DejaVu is the wider of the two, so text that fits here
fits in PowerPoint. Text that overflows here may be fine, which is why the
overflow report prints how much.

    python3 preview_deck.py            # writes preview/slide-N.png
    python3 preview_deck.py --check    # report overflow and margin problems only
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
DECK = HERE / "proposal.pptx"
OUTDIR = HERE / "preview"
SCALE = 110                      # pixels per inch
MARGIN_IN = 0.5                  # the minimum edge margin we hold ourselves to

FONTS = {
    (False, False): "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    (True, False): "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    (False, True): "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
    (True, True): "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
}
_cache: dict = {}


def font(size_pt: float, bold: bool, serif: bool):
    px = max(int(round(size_pt * SCALE / 72.0)), 6)
    key = (px, bold, serif)
    if key not in _cache:
        _cache[key] = ImageFont.truetype(FONTS[(bold, serif)], px)
    return _cache[key]


def emu_px(v) -> float:
    return (v or 0) / 914400.0 * SCALE


def rgb(color, default=(0, 0, 0)):
    try:
        if color and color.type is not None and color.rgb is not None:
            return tuple(color.rgb)
    except Exception:
        pass
    return default


def wrap(draw, text, fnt, width_px):
    """Greedy wrap, the way a text box with word_wrap on behaves."""
    lines, line = [], ""
    for word in text.split():
        trial = f"{line} {word}".strip()
        if draw.textlength(trial, font=fnt) <= width_px or not line:
            line = trial
        else:
            lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def slide_bg(slide):
    try:
        return tuple(slide.background.fill.fore_color.rgb)
    except Exception:
        return (255, 255, 255)


def draw_shape_body(drw, shape, x, y, w, h):
    # An autoshape's shape_type is always AUTO_SHAPE; the geometry is in
    # auto_shape_type. Reading the wrong one draws every disc as a rectangle.
    try:
        name = str(shape.auto_shape_type)
    except Exception:
        name = str(getattr(shape, "shape_type", ""))
    try:
        fill_rgb = tuple(shape.fill.fore_color.rgb)
    except Exception:
        return
    if "OVAL" in name:
        drw.ellipse([x, y, x + w, y + h], fill=fill_rgb)
    elif "ROUNDED" in name:
        drw.rounded_rectangle([x, y, x + w, y + h], radius=min(w, h) * 0.16,
                              fill=fill_rgb)
    elif "ARROW" in name:
        drw.polygon([(x + w * 0.3, y), (x + w * 0.7, y), (x + w * 0.7, y + h * 0.55),
                     (x + w, y + h * 0.55), (x + w / 2, y + h),
                     (x, y + h * 0.55), (x + w * 0.3, y + h * 0.55)],
                    fill=fill_rgb)
    else:
        drw.rectangle([x, y, x + w, y + h], fill=fill_rgb)


def render(prs, check_only=False):
    problems = []
    W = emu_px(prs.slide_width)
    H = emu_px(prs.slide_height)
    OUTDIR.mkdir(exist_ok=True)
    for n, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (int(W), int(H)), slide_bg(slide))
        drw = ImageDraw.Draw(img)
        for shape in slide.shapes:
            x, y = emu_px(shape.left), emu_px(shape.top)
            w, h = emu_px(shape.width), emu_px(shape.height)
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                try:
                    pic = Image.open(shape.image.blob and
                                     __import__("io").BytesIO(shape.image.blob))
                    pic = pic.convert("RGBA").resize((max(int(w), 1), max(int(h), 1)))
                    img.paste(pic, (int(x), int(y)), pic)
                except Exception:
                    drw.rectangle([x, y, x + w, y + h], outline=(200, 0, 0))
                continue
            if not shape.has_text_frame:
                draw_shape_body(drw, shape, x, y, w, h)
                continue
            draw_shape_body(drw, shape, x, y, w, h)

            tf = shape.text_frame
            pad_l = emu_px(tf.margin_left)
            pad_t = emu_px(tf.margin_top)
            inner_w = max(w - pad_l - emu_px(tf.margin_right), 4)
            rendered = []
            for p in tf.paragraphs:
                text = "".join(r.text for r in p.runs)
                if not text:
                    continue
                r0 = p.runs[0]
                size = r0.font.size.pt if r0.font.size else 18
                bold = bool(r0.font.bold)
                serif = (r0.font.name or "").startswith(("Cambria", "Times"))
                fnt = font(size, bold, serif)
                colour = rgb(r0.font.color, (0, 0, 0))
                lead = size * SCALE / 72.0 * (p.line_spacing or 1.18)
                after = (p.space_after.pt if p.space_after else 0) * SCALE / 72.0
                for ln in wrap(drw, text, fnt, inner_w):
                    rendered.append((ln, fnt, colour, lead, str(p.alignment)))
                rendered.append((None, None, None, after, None))

            total = sum(r[3] for r in rendered)
            anchor = str(tf.vertical_anchor)
            cy = y + pad_t
            if "MIDDLE" in anchor:
                cy = y + (h - total) / 2
            for ln, fnt, colour, lead, align in rendered:
                if ln is None:
                    cy += lead
                    continue
                tx = x + pad_l
                if align and "CENTER" in align:
                    tx = x + (w - drw.textlength(ln, font=fnt)) / 2
                elif align and "RIGHT" in align:
                    tx = x + w - pad_l - drw.textlength(ln, font=fnt)
                drw.text((tx, cy), ln, font=fnt, fill=colour)
                cy += lead

            over = (cy - (y + h)) / SCALE
            if over > 0.03:
                problems.append(
                    f"slide {n}: text overflows its box by {over:.2f}in "
                    f"at ({x / SCALE:.2f},{y / SCALE:.2f}) "
                    f"[{''.join(r.text for r in tf.paragraphs[0].runs)[:42]!r}]")
            if x / SCALE < MARGIN_IN - 0.01 or y / SCALE < MARGIN_IN - 0.01 \
                    or (x + w) / SCALE > W / SCALE - MARGIN_IN + 0.01 \
                    or (y + h) / SCALE > H / SCALE - MARGIN_IN + 0.01:
                if shape.has_text_frame and any(r.text for p in tf.paragraphs
                                                for r in p.runs):
                    problems.append(
                        f"slide {n}: box crosses the {MARGIN_IN}in margin at "
                        f"({x / SCALE:.2f},{y / SCALE:.2f}) "
                        f"{w / SCALE:.2f}x{h / SCALE:.2f}")
        if not check_only:
            img.save(OUTDIR / f"slide-{n}.png")
    return problems


def main() -> None:
    prs = Presentation(DECK)
    problems = render(prs, check_only="--check" in sys.argv)
    for p in problems:
        print(p)
    print(f"\n{len(problems)} layout problem(s); "
          f"{len(prs.slides._sldIdLst)} slides")
    if problems:
        sys.exit(1)


if __name__ == "__main__":
    main()
